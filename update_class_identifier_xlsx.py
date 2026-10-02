#!/usr/bin/env python3
"""Update only class identifier cells in a workbook copied from its source."""

import csv
import sys
from pathlib import Path

from openpyxl import load_workbook


def main(source_path: str, output_path: str, mapping_path: str) -> None:
    with Path(mapping_path).open(encoding="utf-8", newline="") as handle:
        mapping = {
            (row["school_id"], row["class_label"]): row["opaque_class"]
            for row in csv.DictReader(handle)
        }

    workbook = load_workbook(source_path)
    if "data" in workbook.sheetnames:
        worksheet = workbook["data"]
    elif len(workbook.sheetnames) == 1:
        worksheet = workbook.active
    else:
        raise ValueError("Workbook must contain a data sheet or exactly one sheet.")
    headers = {
        worksheet.cell(row=1, column=column).value: column
        for column in range(1, worksheet.max_column + 1)
    }
    required = {"school_id", "class_label"}
    if not required.issubset(headers):
        raise ValueError(f"Missing required XLSX columns: {sorted(required - headers.keys())}")

    class_id_column = headers.get("class_id")
    if class_id_column is None:
        class_id_column = worksheet.max_column + 1
        worksheet.cell(row=1, column=class_id_column, value="class_id")

    for row in range(2, worksheet.max_row + 1):
        key = (
            str(worksheet.cell(row=row, column=headers["school_id"]).value),
            str(worksheet.cell(row=row, column=headers["class_label"]).value),
        )
        opaque_class = mapping[key]
        worksheet.cell(row=row, column=headers["class_label"], value=opaque_class)
        worksheet.cell(row=row, column=class_id_column, value=opaque_class)

    workbook.save(output_path)


if __name__ == "__main__":
    if len(sys.argv) != 4:
        raise SystemExit(
            "Usage: update_class_identifier_xlsx.py SOURCE OUTPUT MAPPING_CSV"
        )
    main(*sys.argv[1:])
